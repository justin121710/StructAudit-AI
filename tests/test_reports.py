import io

import openpyxl
import pytest

from agents import schemas as S
from agents.llm_client import MockClient
from orchestrator import ReviewOrchestrator
from utils.parser import parse_structural_file
from utils.report_model import build_model, meta_rows
from utils.xlsx_report import generate_xlsx


@pytest.fixture(scope="module")
def session():
    summary = parse_structural_file("examples/sample_structure.xlsx")
    return ReviewOrchestrator(MockClient()).run(summary, {"案名": "測試案", "設計單位": "某事務所"})


def test_model_header_and_traceability(session):
    m = build_model(session)
    meta = {k: v for k, v, _ in m.meta}
    assert m.project_name == "測試案"
    assert meta["程式版本／雜湊"].startswith("v") and session.source["sha256"][:12] in meta["程式版本／雜湊"]
    # 兩欄式表頭：跨欄項目自成一列，其餘每列兩項
    for row in meta_rows(m.meta):
        assert 1 <= len(row) <= 2
    assert m.sign_roles == ("審查", "校核", "審核")


def test_model_judgement_matches_chief(session):
    m = build_model(session)
    assert len(m.summary_rows) == len(session.issue_ids)
    for row in m.summary_rows:
        assert row.judge == {"Pass": "OK", "Medium": "注意", "High": "NG"}[row.risk]
    assert session.chief.overall_conclusion == S.OverallConclusion.RETURNED
    assert m.verdict_level == "bad" and "NG" in m.verdict_text
    assert len(m.verdict_items) == len(session.chief.major_doubts)


def test_limit_rows_cite_verified_clauses(session):
    refs = " ".join(r for _, _, r in build_model(session).limit_rows)
    assert "§2.16.1" in refs and "§18.4.4.1" in refs and "§5.3.1" in refs


def test_xlsx_layout_and_formula_guard(session):
    session.round1.issues[0].formula = "=HYPERLINK(\"http://x\")"  # 模型輸出以 = 開頭時不得變成公式
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx(session)))
    assert wb.sheetnames == ["檢核表", "審查報告(A4)"]
    for ws in wb:
        assert not [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]
    ws = wb["檢核表"]
    assert ws.freeze_panes == "A3"
    total = [r for r in ws.iter_rows() if r[0].value == "整體審查判定"]
    assert total and total[0][0].fill.fgColor.rgb.endswith("FFF2CC")
    a4 = wb["審查報告(A4)"]
    texts = [c.value for row in a4.iter_rows() for c in row if isinstance(c.value, str)]
    assert "一、審查依據" in texts and "五、檢核彙總" in texts and "審查" in texts
